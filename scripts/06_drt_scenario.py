"""4단계: 관측 보강(탑승 기록) + DRT 시범 시나리오(가정).

A. 관측(2025-05-19~25 1주, 공공데이터포털 '창원시_시내버스 탑승인원')
   - 정류장 코드 체계가 BIS와 달라 '정류장명'으로만 연결할 수 있다.
     그래서 창원 BIS 전체에서 서북동 계곡(Z)에만 존재하는 이름만 사용한다(동명 정류소가 다른 곳에 있으면 제외).
   - 73번 노선의 시간대별 탑승(노선 전체, 계곡 구간만이 아님).
B. 지리(관측): 73번 경유 순서대로 진동환승센터→계곡 정류소까지 정류소 간 직선거리 누적(도로거리보다 짧다).
C. 시나리오(가정): 계곡 ↔ 진동환승센터 호출형 DRT 1대.
   - 운영시간 S1 07~20시, S2 06~23시(가정). 이 시간대에 계곡 주민이 진동에서 탈 수 있는
     '마산역 방향 시각 확인 편'(결과표 1)을 73번 직행과 합쳐 시간대별 '출발 기회 수'를 비교한다.
   - 이것은 연결 가능한 버스 편의 수일 뿐, 실제 이용·대기시간·수요를 뜻하지 않는다.
   - 1대 순환시간 = 왕복 거리 / 가정 평균속도(25·30·35km/h) + 승하차 가정 5분. 용량 상한만 보여준다.

출력: data/processed/boarding_valley.csv, boarding_route73_hourly.csv, drt_geometry.csv,
      drt_scenario_opportunities.csv, drt_scenario_cycle.csv, drt_scenario_summary.json
"""
import json

import numpy as np
import openpyxl
import pandas as pd

from common import OUT, RAW, load_bis, route_stops, route_table, split_directions

BOARD = RAW / "datagokr_cw_bus_boarding_20250519_25.xlsx"
defs = json.loads((OUT / "result_table_1_place_definitions.json").read_text(encoding="utf-8"))
Z_ids = {int(x.split("(")[1].rstrip(")")) for x in defs["Z_서북동계곡"]}

snap = load_bis()
routes = route_table(snap)
rs = route_stops(snap)
names_all = rs.drop_duplicates("spotId")[["spotId", "routeName"]]
z_names = set(names_all.loc[names_all.spotId.isin(Z_ids), "routeName"])
unique_names = sorted(n for n in z_names
                      if set(names_all.loc[names_all.routeName == n, "spotId"]) <= Z_ids)

# ---- A. 탑승 기록 --------------------------------------------------------------
wb = openpyxl.load_workbook(BOARD, read_only=True)
st = pd.DataFrame(list(wb[wb.sheetnames[0]].iter_rows(values_only=True))[1:],
                  columns=["code", "name", "non_transfer", "transfer"])
bv = (st[st.name.isin(unique_names)].groupby("name")[["non_transfer", "transfer"]].sum()
        .assign(total=lambda d: d.non_transfer + d.transfer).reset_index()
        .sort_values("total", ascending=False))
bv.to_csv(OUT / "boarding_valley.csv", index=False, encoding="utf-8-sig")
excluded = sorted(z_names - set(unique_names))

h73 = pd.DataFrame(list(wb["73"].iter_rows(values_only=True))[1:], columns=["hour", "non_transfer", "transfer"])
h73["total"] = h73.non_transfer + h73.transfer
h73.to_csv(OUT / "boarding_route73_hourly.csv", index=False, encoding="utf-8-sig")

# ---- B. 지리 --------------------------------------------------------------------
r73 = split_directions(rs[rs.routeId == "379000730"], routes)
out = r73[r73.updnSe == 0].reset_index(drop=True)
i0 = out.index[out.routeName == "진동환승센터"][0]
seg = pd.concat([out.iloc[i0:], r73[r73.updnSe == 1].head(1)], ignore_index=True)  # 종점(서북동마을) 포함
dx = np.diff(seg.xCrdnt) * 111320 * np.cos(np.radians(35.2))
dy = np.diff(seg.yCrdnt) * 110540
seg["cum_km"] = np.r_[0, np.cumsum(np.hypot(dx, dy))] / 1000
seg["in_valley"] = seg.spotId.isin(Z_ids)
seg[["routeName", "spotId", "cum_km", "in_valley"]].to_csv(OUT / "drt_geometry.csv", index=False, encoding="utf-8-sig")
d_first = float(seg.loc[seg.in_valley, "cum_km"].min())
d_last = float(seg.cum_km.max())

# ---- C. 시나리오 ------------------------------------------------------------------
long = pd.read_csv(OUT / "result_table_1_long.csv", encoding="utf-8-sig")
BANDS = ["~07시", "07–09시", "09–12시", "12–15시", "15–18시", "18–20시", "20시~"]
BAND_HOURS = {"~07시": (0, 7), "07–09시": (7, 9), "09–12시": (9, 12), "12–15시": (12, 15),
              "15–18시": (15, 18), "18–20시": (18, 20), "20시~": (20, 24)}
SCEN = {"S1 07~20시": (7, 20), "S2 06~23시": (6, 23)}
rows = []
for day in ["평일", "일·공휴일"]:
    for direction, zdir, cdir in [("마산역 방향", "생활권→마산역", "생활권→마산역"),
                                  ("귀가 방향", "마산역→생활권", "마산역→생활권")]:
        z = long[(long["구분"].str.startswith("대상")) & (long["방향"] == zdir) & (long["요일"] == day)]
        c = long[(long["구분"].str.startswith("비교")) & (long["방향"] == cdir) & (long["요일"] == day)
                 & (long["시간대"] != "시각 미산출(중간경유)")]
        # 73번은 진동도 지나므로 진동 쪽 편에서 73번을 빼 중복을 막는다
        c = c[c["노선"].astype(str) != "73"]
        for sname, (a, b) in SCEN.items():
            rec = {"요일": day, "방향": direction, "시나리오": sname}
            for band in BANDS:
                lo, hi = BAND_HOURS[band]
                base = int((z["시간대"] == band).sum())
                mins = c["출발시각"].str.slice(0, 2).astype(int) * 60 + c["출발시각"].str.slice(3, 5).astype(int)
                in_band = c[(mins >= lo * 60) & (mins < hi * 60) & (mins >= a * 60) & (mins < b * 60)]
                rec[f"{band}_현행"] = base
                rec[f"{band}_시나리오"] = base + len(in_band)
            rec["합계_현행"] = sum(rec[f"{bd}_현행"] for bd in BANDS)
            rec["합계_시나리오"] = sum(rec[f"{bd}_시나리오"] for bd in BANDS)
            rows.append(rec)
opp = pd.DataFrame(rows)
opp.to_csv(OUT / "drt_scenario_opportunities.csv", index=False, encoding="utf-8-sig")

cyc = []
for v in [25, 30, 35]:
    for dwell in [5]:
        m = 2 * d_last / v * 60 + dwell
        cyc.append({"가정 평균속도(km/h)": v, "가정 승하차·회차(분)": dwell,
                    "진동↔서북동마을 왕복(분)": round(m), "1대 시간당 최대 왕복": round(60 / m, 2)})
pd.DataFrame(cyc).to_csv(OUT / "drt_scenario_cycle.csv", index=False, encoding="utf-8-sig")

summary = {
    "valley_unique_name_stops": unique_names, "valley_names_excluded_ambiguous": excluded,
    "valley_weekly_boardings_unique_names": int(bv.total.sum()),
    "route73_weekly_boardings_all_stops": int(h73.total.sum()),
    "dist_km_jindong_to_first_valley_stop": round(d_first, 2),
    "dist_km_jindong_to_seobukdong_terminal": round(d_last, 2),
}
(OUT / "drt_scenario_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=2))
print(bv.to_string(index=False))
print(opp.to_string(index=False))
print(pd.DataFrame(cyc).to_string(index=False))
