"""1단계: 원본 파일 목록(크기·SHA-256·행·열)을 기록한다. 원본은 수정하지 않는다.
출력: data/processed/source_inventory.csv
"""
import hashlib
import json

import openpyxl
import pandas as pd

from common import OUT, RAW

rows = []
for p in sorted(RAW.rglob("*")):
    if not p.is_file() or p.name.startswith("."):
        continue
    info = {"file": str(p.relative_to(RAW)), "bytes": p.stat().st_size,
            "sha256": hashlib.sha256(p.read_bytes()).hexdigest()[:16], "rows": None, "columns": None}
    if p.suffix == ".csv":
        df = pd.read_csv(p, encoding="utf-8-sig")
        info.update(rows=len(df), columns="|".join(df.columns))
    elif p.suffix == ".xlsx":
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        info.update(rows=sum(wb[s].max_row or 0 for s in wb.sheetnames),
                    columns="sheets: " + "|".join(wb.sheetnames))
    elif p.suffix == ".json":
        j = json.loads(p.read_text(encoding="utf-8"))
        info.update(rows=f"routes={len(j['route_list'])}, route_stop_rows={sum(len(v) for v in j['route_stops'].values())}, "
                         f"timetable_rows={sum(len(v or []) for v in j['timetables'].values())}",
                    columns="keys: " + "|".join(j.keys()))
    rows.append(info)

inv = pd.DataFrame(rows)
inv.to_csv(OUT / "source_inventory.csv", index=False, encoding="utf-8-sig")
print(inv.to_string(index=False))
